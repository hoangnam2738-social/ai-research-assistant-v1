import streamlit as st
import pandas as pd
import pymupdf
import docx
from docx.shared import Pt, Inches, RGBColor # Thêm RGBColor để chỉnh màu chữ
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
import os
from dotenv import load_dotenv
from google import genai
import matplotlib.pyplot as plt
import seaborn as sns

# --- KHẮC PHỤC 1: ÉP MATPLOTLIB DÙNG FONT HỖ TRỢ TIẾNG VIỆT ---
plt.rcParams['font.family'] = 'sans-serif'
plt.rcParams['font.sans-serif'] = ['Arial', 'Tahoma', 'Segoe UI', 'DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False

# 1. CẤU HÌNH API
load_dotenv()
# Thử lấy khóa từ Streamlit Cloud, nếu chạy ở máy tính local thì lấy từ file .env
try:
    api_key = st.secrets["GEMINI_API_KEY"]
except:
    api_key = os.getenv("GEMINI_API_KEY")

client = genai.Client(api_key=api_key)

# 2. CÁC HÀM XỬ LÝ LÕI
def extract_text(uploaded_file):
    try:
        doc = pymupdf.open(stream=uploaded_file.read(), filetype="pdf")
        text = "".join([page.get_text() for page in doc])
        return text[:15000] 
    except:
        return ""

def generate_section(prompt):
    response = client.models.generate_content(
        model='gemini-3.5-flash',
        contents=prompt
    )
    return response.text

def add_formatted_text(document, text):
    for line in text.split('\n'):
        line = line.strip()
        if not line:
            continue
            
        # Tự động bộ lọc xóa các dấu markdown và latex thô
        clean_line = line.replace('**', '').replace('$', '').replace('`', '') 
        
        p = document.add_paragraph(clean_line)
        
        if line.startswith('#'):
            p.runs[0].bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5

# 3. GIAO DIỆN CHÍNH
st.set_page_config(page_title="AI Research Engine V3.1", layout="wide")
st.title("🚀 Cỗ Máy Sản Xuất Bài Báo Học Thuật V3.1 (Đã fix Font)")

topic = st.text_input("Tên đề tài nghiên cứu:", "Nghịch lý của sự lựa chọn trong AI Recommendations trên E-commerce")

col1, col2 = st.columns(2)
with col1:
    st.subheader("1. Nạp Tài liệu (Literature Review)")
    uploaded_pdfs = st.file_uploader("Tải lên các file PDF (bài báo tham khảo)", type="pdf", accept_multiple_files=True)

with col2:
    st.subheader("2. Nạp Dữ liệu (Methodology & Results)")
    uploaded_csv = st.file_uploader("Tải lên file dữ liệu hành vi (CSV)", type="csv")

if st.button("🚀 Xử lý Dữ liệu, Vẽ Biểu Đồ & Viết Bài Báo", type="primary"):
    
    # --- KHẮC PHỤC 2: ĐỒNG BỘ TOÀN BỘ FONT TRONG FILE WORD ---
    doc = docx.Document()
    
    # Chỉnh Font Normal (Văn bản thường)
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    
    # Chỉnh Font Heading 1 (Tiêu đề Chương: Chuyển sang Đen, Times New Roman, Bold)
    style_h1 = doc.styles['Heading 1']
    style_h1.font.name = 'Times New Roman'
    style_h1.font.size = Pt(14)
    style_h1.font.bold = True
    style_h1.font.color.rgb = RGBColor(0, 0, 0)
    
    # Chỉnh Font Title (Tiêu đề chính)
    style_title = doc.styles['Title']
    style_title.font.name = 'Times New Roman'
    style_title.font.size = Pt(16)
    style_title.font.bold = True
    style_title.font.color.rgb = RGBColor(0, 0, 0)
    
    # Tạo Trang bìa
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.add_run("\n\n\n\nBÁO CÁO NGHIÊN CỨU KHOA HỌC\n").bold = True
    doc.add_heading(topic.upper(), level=0) # Dùng Heading 0 (Title)
    cover_sub = doc.add_paragraph()
    cover_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_sub.add_run("\nĐược tạo tự động bởi AI Research Engine")
    doc.add_page_break()
    
    # --- PHASE A: ĐỌC VÀ TỔNG HỢP PDF ---
    lit_review_knowledge = ""
    if uploaded_pdfs:
        with st.status("Đang đọc và phân tích PDF..."):
            for pdf in uploaded_pdfs:
                text = extract_text(pdf)
                lit_review_knowledge += f"\nNội dung từ {pdf.name}: {text}\n"
    else:
        lit_review_knowledge = "Không có tài liệu tham khảo nào được cung cấp."

    # --- PHASE B: PHÂN TÍCH DATA & VẼ BIỂU ĐỒ ---
    data_summary = ""
    img_buffer = BytesIO() 
    
    if uploaded_csv:
        with st.status("Đang phân tích số liệu và vẽ biểu đồ trực quan..."):
            df = pd.read_csv(uploaded_csv)
            bins = [0, 5, 10, 15, 20]
            labels = ['0-5', '6-10', '11-15', '16-20']
            df['Group'] = pd.cut(df['RecNum'], bins=bins, labels=labels, include_lowest=True)
            summary = df.groupby('Group', observed=False)['Abandon'].mean() * 100
            
            data_summary = "Kết quả chạy phân tích định lượng:\n"
            for g, r in summary.items():
                data_summary += f"- Nhóm gợi ý {g} sản phẩm: Tỷ lệ bỏ giỏ = {r:.2f}%\n"
            
            # Vẽ biểu đồ
            sns.set_theme(style="whitegrid")
            plt.rcParams['font.family'] = 'sans-serif' # Nhắc lại thiết lập font cho chắc chắn
            plt.rcParams['font.sans-serif'] = ['Arial', 'Tahoma', 'Segoe UI']
            
            fig, ax = plt.subplots(figsize=(8, 5))
            ax.plot(labels, summary.values, marker='o', color='#e74c3c', linestyle='-', linewidth=2.5, markersize=8)
            ax.set_title('Tác động của Số lượng Gợi ý AI đến Tỷ lệ Rời bỏ', fontsize=14, fontweight='bold')
            ax.set_xlabel('Số lượng Gợi ý (Sản phẩm)', fontsize=12)
            ax.set_ylabel('Tỷ lệ Rời bỏ Giỏ hàng (%)', fontsize=12)
            
            fig.savefig(img_buffer, format='png', bbox_inches='tight')
            img_buffer.seek(0)

    # --- PHASE C: GỌI AI & HIỂN THỊ LIVE PREVIEW ---
    # --- PHASE C: GỌI AI & HIỂN THỊ LIVE PREVIEW ---
    with st.spinner("AI đang chắp bút viết bài... (Vui lòng chờ khoảng 45s)"):
        
        prompt_1 = f"""
        Đóng vai Giáo sư Kinh tế học Hành vi và Khoa học Dữ liệu. Viết Chương 1 và Chương 2 cho đề tài '{topic}'.
        Nguồn tài liệu (BẮT BUỘC trích dẫn): {lit_review_knowledge}
        
        YÊU CẦU NỘI DUNG: Trích dẫn chuẩn APA. Xây dựng Khung lý thuyết và Khoảng trống nghiên cứu. Hành văn súc tích. Viết khoảng 800 từ.
        
        YÊU CẦU ĐỊNH DẠNG (BẮT BUỘC TUÂN THỦ NGHIÊM NGẶT):
        - KHÔNG vẽ sơ đồ bằng ký tự (như +---+ hoặc |). Hãy mô tả khung lý thuyết bằng các đoạn văn.
        - KHÔNG dùng định dạng bảng Markdown (như | Cột 1 | Cột 2 |). Trình bày thông tin dưới dạng danh sách gạch đầu dòng.
        """
        text_1 = generate_section(prompt_1)
        doc.add_heading('1. Giới thiệu & Tổng quan tài liệu', level=1)
        add_formatted_text(doc, text_1)

        prompt_2 = f"""
        Viết Chương 3 (Phương pháp luận) và Chương 4 (Phân tích kết quả) cho đề tài '{topic}'.
        DỮ LIỆU ĐẦU VÀO: {data_summary}
        
        YÊU CẦU NỘI DUNG: Thiết lập mô hình Hồi quy Logistic. Biện luận ý nghĩa thống kê. Giải thích 'điểm gãy' bằng giới hạn nhận thức. Viết khoảng 1000 từ.
        
        YÊU CẦU ĐỊNH DẠNG (BẮT BUỘC TUÂN THỦ NGHIÊM NGẶT):
        - KHÔNG dùng bất kỳ ký hiệu toán học LaTeX nào (TUYỆT ĐỐI KHÔNG dùng dấu $, $$, \beta, \frac, \epsilon). Hãy viết tên các biến và công thức bằng chữ thuần túy (Ví dụ: Beta 1, P(Y=1), Xác suất bỏ giỏ = ...).
        - KHÔNG dùng định dạng bảng Markdown (như | Nhóm | Tỷ lệ |). Hãy liệt kê kết quả số liệu thành các gạch đầu dòng rõ ràng.
        """
        text_2 = generate_section(prompt_2)
        doc.add_heading('2. Phương pháp và Kết quả Định lượng', level=1)
        
        if uploaded_csv:
            doc.add_paragraph("Biểu đồ dưới đây minh họa trực quan sự thay đổi tỷ lệ rời bỏ giỏ hàng theo số lượng gợi ý, xác nhận giả thuyết về cấu trúc hình chữ U ngược của nghịch lý lựa chọn.")
            doc.add_picture(img_buffer, width=Inches(5.5))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_caption = doc.add_paragraph("Hình 1: Tác động của số lượng gợi ý đến tỷ lệ rời bỏ giỏ hàng")
            img_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_caption.style.font.italic = True
        
        add_formatted_text(doc, text_2)

    st.success("🎉 HOÀN THÀNH TOÀN BỘ BÀI BÁO!")
    
    with st.expander("👀 XEM TRƯỚC BẢN THẢO (LIVE PREVIEW)", expanded=True):
        st.markdown("### Chương 1 & 2: Giới thiệu & Tổng quan tài liệu")
        st.write(text_1)
        
        st.markdown("### Chương 3 & 4: Phương pháp & Kết quả")
        if uploaded_csv:
            st.image(img_buffer, caption="Hình 1: Tác động của số lượng gợi ý đến tỷ lệ rời bỏ giỏ hàng")
        st.write(text_2)

    # --- PHASE D: XUẤT FILE ---
    bio = BytesIO()
    doc.save(bio)
    
    st.download_button(
        label="📥 Tải Bài Báo Hoàn Chỉnh Về Máy (.docx)",
        data=bio.getvalue(),
        file_name="Ultimate_Research_Paper_FixedFont.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )